"""Per-document transfer report: job-level totals + row-per-document details."""
from __future__ import annotations

import io
import re
import threading
import zipfile
from datetime import datetime
from pathlib import Path
from typing import BinaryIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .document_filter import is_document, load_document_extensions
from .excel_io import load_workbook_with_retry, save_workbook_with_retry
from .logging_setup import get_logger
from .path_utils import iter_files


_JOBS_SHEET = "_Jobs"
_JOBS_HEADERS = [
    "Job ID",
    "Started at",
    "Finished at",
    "Study",
    "Source name",
    "Source type",
    "Source docs",
    "Destination docs",
    "Docs match",
    "Source size (bytes)",
    "Destination size (bytes)",
    "Size match",
    "Status",
    "Notes",
]

_DOC_HEADERS = [
    "Job ID",
    "Timestamp",
    "Relative path",
    "Size (bytes)",
    "Origin",
    "Status",
]

_INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


class TransferReporter:
    """xlsx workbook with a _Jobs summary sheet and per-study document sheets."""

    def __init__(self, report_path: Path) -> None:
        self.report_path = report_path
        self._lock = threading.Lock()
        self._log = get_logger()
        self.report_path.parent.mkdir(parents=True, exist_ok=True)
        self._ensure_workbook()

    # ------------------------------------------------------------------ helpers

    @staticmethod
    def _now() -> str:
        return datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    @staticmethod
    def _safe_sheet_name(name: str) -> str:
        cleaned = _INVALID_SHEET_CHARS.sub("_", name).strip() or "Study"
        return cleaned[:31]

    @staticmethod
    def zip_totals(zip_path: Path) -> tuple[int, int, list[tuple[str, int]]]:
        """(doc_count, total_uncompressed_size, [(name, size), ...]) — recurses nested zips, whitelist-filtered."""
        docs: list[tuple[str, int]] = []
        allowed = load_document_extensions()
        try:
            with zipfile.ZipFile(zip_path, "r") as z:
                TransferReporter._collect_zip_entries(z, prefix="", out=docs, allowed=allowed)
        except (zipfile.BadZipFile, OSError) as exc:
            get_logger().warning("Failed to inspect ZIP %s: %s", zip_path, exc)
            return 0, 0, []
        return len(docs), sum(size for _, size in docs), docs

    @staticmethod
    def _collect_zip_entries(
        zf: zipfile.ZipFile,
        prefix: str,
        out: list[tuple[str, int]],
        allowed: frozenset[str],
    ) -> None:
        for info in zf.infolist():
            if info.is_dir():
                continue
            name = f"{prefix}{info.filename}"
            if name.lower().endswith(".zip"):
                try:
                    with zf.open(info, "r") as raw:
                        data = raw.read()
                    TransferReporter._collect_from_bytes(data, f"{name}/", out, allowed)
                except (zipfile.BadZipFile, OSError) as exc:
                    get_logger().warning("Failed to inspect nested ZIP %s: %s", name, exc)
                continue
            if is_document(name, allowed):
                out.append((name, info.file_size))

    @staticmethod
    def _collect_from_bytes(
        data: bytes, prefix: str, out: list[tuple[str, int]], allowed: frozenset[str]
    ) -> None:
        buffer: BinaryIO = io.BytesIO(data)
        with zipfile.ZipFile(buffer, "r") as nested:
            TransferReporter._collect_zip_entries(nested, prefix, out, allowed)

    @staticmethod
    def folder_totals(folder: Path) -> tuple[int, int, list[tuple[Path, int]]]:
        """(doc_count, total_size, [(path, size), ...]) — whitelist-filtered, long-path safe."""
        if not folder.exists():
            return 0, 0, []
        allowed = load_document_extensions()
        docs: list[tuple[Path, int]] = []
        total_size = 0
        for path, size in iter_files(folder):
            if not is_document(path.name, allowed):
                continue
            docs.append((path, size))
            total_size += size
        return len(docs), total_size, docs

    # ---------------------------------------------------------------- public API

    def record_zip_job(
        self,
        study_name: str,
        zip_path: Path,
        extracted_dir: Path,
        status: str,
        notes: str = "",
        started_at: str | None = None,
    ) -> str:
        src_count, src_size, _src_docs = self.zip_totals(zip_path)
        dest_count, dest_size, dest_docs = self.folder_totals(extracted_dir)
        finished = self._now()
        started = started_at or finished

        with self._lock:
            wb = load_workbook_with_retry(self.report_path)
            jobs = self._get_or_create_jobs_sheet(wb)
            job_id = self._next_job_id(jobs)
            jobs.append(
                [
                    job_id,
                    started,
                    finished,
                    study_name,
                    zip_path.name,
                    "ZIP",
                    src_count,
                    dest_count,
                    src_count == dest_count,
                    src_size,
                    dest_size,
                    src_size == dest_size,
                    status,
                    notes,
                ]
            )
            docs_sheet = self._get_or_create_docs_sheet(wb, study_name)
            for path, size in dest_docs:
                try:
                    rel = path.relative_to(extracted_dir)
                except ValueError:
                    rel = path
                docs_sheet.append([job_id, finished, str(rel), size, "ZIP", "TRANSFERRED"])
            self._autosize(jobs)
            self._autosize(docs_sheet)
            save_workbook_with_retry(wb, self.report_path)

        self._log.info(
            "Transfer report: job=%s study=%s src_docs=%d dest_docs=%d src_size=%d dest_size=%d",
            job_id,
            study_name,
            src_count,
            dest_count,
            src_size,
            dest_size,
        )
        return job_id

    def record_direct_transfer(
        self, source_path: Path, dest_path: Path, status: str = "TRANSFERRED", notes: str = ""
    ) -> str:
        try:
            src_size = source_path.stat().st_size if source_path.exists() else dest_path.stat().st_size
        except OSError:
            src_size = 0
        try:
            dest_size = dest_path.stat().st_size
        except OSError:
            dest_size = 0
        finished = self._now()

        with self._lock:
            wb = load_workbook_with_retry(self.report_path)
            jobs = self._get_or_create_jobs_sheet(wb)
            job_id = self._next_job_id(jobs)
            jobs.append(
                [
                    job_id,
                    finished,
                    finished,
                    "_Direct",
                    source_path.name,
                    "FILE",
                    1,
                    1 if dest_path.exists() else 0,
                    dest_path.exists(),
                    src_size,
                    dest_size,
                    src_size == dest_size,
                    status,
                    notes,
                ]
            )
            docs_sheet = self._get_or_create_docs_sheet(wb, "_Direct")
            docs_sheet.append([job_id, finished, dest_path.name, dest_size, "FILE", status])
            self._autosize(jobs)
            self._autosize(docs_sheet)
            save_workbook_with_retry(wb, self.report_path)
        return job_id

    # -------------------------------------------------------------- workbook ops

    def _ensure_workbook(self) -> None:
        if self.report_path.exists():
            return
        wb = Workbook()
        default = wb.active
        if default is not None:
            wb.remove(default)
        self._get_or_create_jobs_sheet(wb)
        save_workbook_with_retry(wb, self.report_path)

    def _get_or_create_jobs_sheet(self, wb: Workbook):
        if _JOBS_SHEET in wb.sheetnames:
            return wb[_JOBS_SHEET]
        ws = wb.create_sheet(title=_JOBS_SHEET, index=0)
        ws.append(_JOBS_HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        return ws

    def _get_or_create_docs_sheet(self, wb: Workbook, study_name: str):
        title = self._safe_sheet_name(study_name)
        if title == _JOBS_SHEET:
            title = f"{title}_docs"[:31]
        if title in wb.sheetnames:
            return wb[title]
        ws = wb.create_sheet(title=title)
        ws.append(_DOC_HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        return ws

    @staticmethod
    def _next_job_id(jobs_sheet) -> str:
        next_number = max(1, jobs_sheet.max_row)  # header counts as row 1
        return f"J{next_number:06d}"

    @staticmethod
    def _autosize(ws) -> None:
        for column_cells in ws.columns:
            length = 12
            letter = get_column_letter(column_cells[0].column)
            for cell in column_cells:
                value = cell.value
                if value is None:
                    continue
                length = max(length, min(60, len(str(value)) + 2))
            ws.column_dimensions[letter].width = length
