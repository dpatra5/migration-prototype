"""Excel reporter: per-study sheet with document counts before/after extraction."""
from __future__ import annotations

import re
import threading
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .document_filter import is_document, load_document_extensions
from .excel_io import load_workbook_with_retry, save_workbook_with_retry
from .logging_setup import get_logger


_EVENT_HEADERS = [
    "Timestamp",
    "ZIP filename",
    "Extracted folder",
    "Docs before extraction",
    "Docs after extraction",
    "Status",
    "Notes",
]

_SNAPSHOT_HEADERS = ["Timestamp", "Study folder", "Current document count"]
_SNAPSHOT_SHEET = "_Snapshot"

_INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


class ExcelReporter:
    """Maintains an xlsx workbook with one sheet per study plus a periodic snapshot sheet."""

    def __init__(self, report_path: Path, destination_root: Path) -> None:
        self.report_path = report_path
        self.destination_root = destination_root
        self._lock = threading.Lock()
        self._log = get_logger()
        self.report_path.parent.mkdir(parents=True, exist_ok=True)
        self._ensure_workbook()

    # ------------------------------------------------------------------ counts

    @staticmethod
    def count_zip_documents(zip_path: Path) -> int:
        """Count whitelisted document entries inside a ZIP (nested ZIPs expanded)."""
        allowed = load_document_extensions()

        def _walk(zf: zipfile.ZipFile, prefix: str) -> int:
            total = 0
            for info in zf.infolist():
                if info.is_dir():
                    continue
                name = f"{prefix}{info.filename}"
                if name.lower().endswith(".zip"):
                    try:
                        with zf.open(info, "r") as raw:
                            data = raw.read()
                        import io as _io
                        with zipfile.ZipFile(_io.BytesIO(data), "r") as nested:
                            total += _walk(nested, f"{name}/")
                    except (zipfile.BadZipFile, OSError):
                        continue
                elif is_document(name, allowed):
                    total += 1
            return total

        try:
            with zipfile.ZipFile(zip_path, "r") as z:
                return _walk(z, "")
        except (zipfile.BadZipFile, OSError):
            return 0

    @staticmethod
    def count_folder_documents(folder: Path) -> int:
        if not folder.exists():
            return 0
        allowed = load_document_extensions()
        from .path_utils import iter_files
        return sum(1 for path, _size in iter_files(folder) if is_document(path.name, allowed))

    # ---------------------------------------------------------------- public API

    def record_extraction(
        self,
        study_name: str,
        zip_path: Path,
        extracted_folder: Path,
        docs_before: int,
        docs_after: int,
        status: str = "COMPLETED",
        notes: str = "",
    ) -> None:
        """Append a row to the study's sheet describing this extraction."""
        with self._lock:
            wb = load_workbook_with_retry(self.report_path)
            ws = self._get_or_create_event_sheet(wb, study_name)
            ws.append(
                [
                    datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    zip_path.name,
                    str(extracted_folder),
                    docs_before,
                    docs_after,
                    status,
                    notes,
                ]
            )
            self._autosize(ws)
            save_workbook_with_retry(wb, self.report_path)
        self._log.info(
            "Reported extraction: study=%s before=%d after=%d", study_name, docs_before, docs_after
        )

    def refresh_snapshot(self) -> None:
        """Re-scan DESTINATION_FOLDER and write a fresh snapshot row per study."""
        if not self.destination_root.exists():
            self._log.warning("Destination folder missing; snapshot skipped: %s", self.destination_root)
            return

        studies = [p for p in self.destination_root.iterdir() if p.is_dir()]
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with self._lock:
            wb = load_workbook_with_retry(self.report_path)
            ws = self._get_or_create_snapshot_sheet(wb)
            for study_dir in sorted(studies, key=lambda p: p.name.lower()):
                count = self.count_folder_documents(study_dir)
                ws.append([timestamp, study_dir.name, count])
            self._autosize(ws)
            save_workbook_with_retry(wb, self.report_path)
        self._log.info("Snapshot refreshed for %d studies", len(studies))

    # -------------------------------------------------------------- workbook ops

    def _ensure_workbook(self) -> None:
        if self.report_path.exists():
            return
        wb = Workbook()
        default = wb.active
        if default is not None:
            wb.remove(default)
        self._get_or_create_snapshot_sheet(wb)
        save_workbook_with_retry(wb, self.report_path)

    def _get_or_create_event_sheet(self, wb: Workbook, study_name: str):
        title = self._safe_sheet_name(study_name)
        if title in wb.sheetnames:
            return wb[title]
        ws = wb.create_sheet(title=title)
        ws.append(_EVENT_HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        return ws

    def _get_or_create_snapshot_sheet(self, wb: Workbook):
        if _SNAPSHOT_SHEET in wb.sheetnames:
            return wb[_SNAPSHOT_SHEET]
        ws = wb.create_sheet(title=_SNAPSHOT_SHEET, index=0)
        ws.append(_SNAPSHOT_HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        return ws

    @staticmethod
    def _safe_sheet_name(name: str) -> str:
        cleaned = _INVALID_SHEET_CHARS.sub("_", name).strip() or "Study"
        return cleaned[:31]

    @staticmethod
    def _autosize(ws) -> None:
        for column_cells in ws.columns:
            length = 12
            letter = get_column_letter(column_cells[0].column)
            for cell in column_cells:
                value = cell.value
                if value is None:
                    continue
                length = max(length, min(60, len(str(value)) + 2))
            ws.column_dimensions[letter].width = length
