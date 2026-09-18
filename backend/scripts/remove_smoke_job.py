"""One-shot cleanup: strip the StudyX smoke-test job (originally J000001 / J-001)
from every transfer_report workbook (primary + pending sidecars).

Deletes the matching row from the `_Jobs` sheet and drops the `StudyX` docs
sheet if present. Creates a `.bak` copy of each file before editing.
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook


REPORT_DIR = Path(r"C:\Users\ADas155\QuIn\POC\Extracted_Folders")
JOBS_SHEET = "_Jobs"
SMOKE_STUDY = "StudyX"
SMOKE_JOB_IDS = {"J000001", "J-001"}


def clean_workbook(path: Path) -> tuple[int, bool]:
    """Return (rows_removed, docs_sheet_removed)."""
    wb = load_workbook(path)
    removed = 0
    docs_removed = False

    if JOBS_SHEET in wb.sheetnames:
        ws = wb[JOBS_SHEET]
        # openpyxl rows are 1-indexed; row 1 is header.
        for row_idx in range(ws.max_row, 1, -1):
            job_id = str(ws.cell(row=row_idx, column=1).value or "").strip()
            study = str(ws.cell(row=row_idx, column=4).value or "").strip()
            if job_id in SMOKE_JOB_IDS or study.casefold() == SMOKE_STUDY.casefold():
                ws.delete_rows(row_idx, 1)
                removed += 1

    if SMOKE_STUDY in wb.sheetnames:
        del wb[SMOKE_STUDY]
        docs_removed = True

    if removed or docs_removed:
        backup = path.with_suffix(path.suffix + ".bak")
        if not backup.exists():
            shutil.copy2(path, backup)
        wb.save(path)
    wb.close()
    return removed, docs_removed


def main() -> int:
    targets = sorted(REPORT_DIR.glob("transfer_report*.xlsx"))
    if not targets:
        print(f"No transfer_report workbooks under {REPORT_DIR}")
        return 1
    for path in targets:
        try:
            removed, docs_removed = clean_workbook(path)
        except PermissionError as exc:
            print(f"SKIP {path.name}: locked ({exc}). Close it in Excel and rerun.")
            continue
        print(
            f"{path.name}: removed {removed} row(s)"
            f"{' + StudyX docs sheet' if docs_removed else ''}"
        )
    return 0


if __name__ == "__main__":
    sys.exit(main())
