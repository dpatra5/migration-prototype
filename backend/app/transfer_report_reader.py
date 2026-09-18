"""Read the extractor's `transfer_report.xlsx` so the dashboard can surface
job history that predates (or lives outside) the SQLite Jobs table.

The workbook layout is produced by `zip_extractor.transfer_report.TransferReporter`:

- `_Jobs` sheet columns (in order):
    Job ID, Started at, Finished at, Study, Source name, Source type,
    Source docs, Destination docs, Docs match, Source size (bytes),
    Destination size (bytes), Size match, Status, Notes,
    Mapped docs, Unclassified docs
- One sheet per study with per-document detail rows.
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Iterable

try:
    from openpyxl import load_workbook  # type: ignore
except Exception:  # pragma: no cover - openpyxl is a required dep in requirements.txt
    load_workbook = None  # type: ignore


_JOBS_SHEET = "_Jobs"
_DEFAULT_REPORT = Path(r"C:\Users\ADas155\QuIn\POC\Extracted_Folders\transfer_report.xlsx")

# Legacy writer used "J000006"; new format is "J-006". Normalize on read.
_LEGACY_JOB_ID_RE = re.compile(r"^J(\d{4,})$")


def report_path() -> Path:
    return Path(os.getenv("TRANSFER_REPORT_PATH", str(_DEFAULT_REPORT)))


def _pick_active_report(path: Path) -> Path:
    """Return the freshest workbook: the primary or a `.pending-*` sidecar.

    When the primary is locked (open in Excel), the writer saves a
    `<stem>.pending-<timestamp><suffix>` sidecar. That sidecar holds the latest
    mapped / unclassified counts, so the dashboard should read whichever file
    has the newest mtime.
    """
    parent = path.parent
    stem = path.stem
    suffix = path.suffix
    candidates: list[Path] = []
    if path.exists():
        candidates.append(path)
    if parent.exists():
        candidates.extend(parent.glob(f"{stem}.pending-*{suffix}"))
    if not candidates:
        return path
    return max(candidates, key=lambda p: p.stat().st_mtime)


def _normalize_job_id(raw: str) -> str:
    m = _LEGACY_JOB_ID_RE.match(raw.strip())
    if not m:
        return raw
    return f"J-{int(m.group(1)):03d}"


@dataclass
class TransferJobRow:
    jobId: str
    startedAt: str | None
    finishedAt: str | None
    study: str
    sourceName: str
    sourceType: str
    sourceDocs: int
    destinationDocs: int
    docsMatch: bool
    sourceSize: int
    destinationSize: int
    sizeMatch: bool
    status: str
    notes: str
    mappedDocs: int
    unclassifiedDocs: int


@dataclass
class TransferReportSummary:
    totalJobs: int = 0
    totalMapped: int = 0
    totalUnclassified: int = 0
    totalSourceDocs: int = 0
    totalDestinationDocs: int = 0
    statuses: dict[str, int] = field(default_factory=dict)
    studies: set[str] = field(default_factory=set)


def _as_int(value) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return 0


def _as_bool(value) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    text = str(value).strip().lower()
    return text in {"true", "1", "yes"}


def _as_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value)


def read_jobs(path: Path | None = None) -> list[TransferJobRow]:
    """Return every row from the `_Jobs` sheet, newest last (as recorded).

    Automatically picks the freshest workbook between the primary
    `transfer_report.xlsx` and any `transfer_report.pending-*.xlsx` sidecars,
    so mapped / unclassified counts reflect the latest run even when the
    primary was locked (open in Excel) at save time.
    """
    target = _pick_active_report(path or report_path())
    if load_workbook is None or not target.exists():
        return []
    try:
        wb = load_workbook(target, read_only=True, data_only=True)
    except Exception:
        return []
    try:
        if _JOBS_SHEET not in wb.sheetnames:
            return []
        ws = wb[_JOBS_SHEET]
        rows_iter = ws.iter_rows(values_only=True)
        headers = next(rows_iter, None)
        if not headers:
            return []
        jobs: list[TransferJobRow] = []
        for row in rows_iter:
            if not row or all(cell is None for cell in row):
                continue
            padded = list(row) + [None] * (16 - len(row))
            (
                job_id, started, finished, study, source_name, source_type,
                src_docs, dst_docs, docs_match, src_size, dst_size,
                size_match, status, notes, mapped, unclass,
            ) = padded[:16]
            if not job_id:
                continue
            jobs.append(
                TransferJobRow(
                    jobId=_normalize_job_id(_as_text(job_id)),
                    startedAt=_as_text(started) or None,
                    finishedAt=_as_text(finished) or None,
                    study=_as_text(study),
                    sourceName=_as_text(source_name),
                    sourceType=_as_text(source_type),
                    sourceDocs=_as_int(src_docs),
                    destinationDocs=_as_int(dst_docs),
                    docsMatch=_as_bool(docs_match),
                    sourceSize=_as_int(src_size),
                    destinationSize=_as_int(dst_size),
                    sizeMatch=_as_bool(size_match),
                    status=_as_text(status) or "UNKNOWN",
                    notes=_as_text(notes),
                    mappedDocs=_as_int(mapped),
                    unclassifiedDocs=_as_int(unclass),
                )
            )
        return jobs
    finally:
        wb.close()


def summarize(jobs: Iterable[TransferJobRow]) -> TransferReportSummary:
    summary = TransferReportSummary()
    for job in jobs:
        summary.totalJobs += 1
        summary.totalMapped += job.mappedDocs
        summary.totalUnclassified += job.unclassifiedDocs
        summary.totalSourceDocs += job.sourceDocs
        summary.totalDestinationDocs += job.destinationDocs
        key = job.status.upper() or "UNKNOWN"
        summary.statuses[key] = summary.statuses.get(key, 0) + 1
        if job.study:
            summary.studies.add(job.study)
    return summary
